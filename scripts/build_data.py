"""
Bake the Japan Tourism Agency prefecture data + map boundaries into static JSON
for the dashboard. Run manually when a new annual JTA workbook is published:

    /usr/local/bin/python3 scripts/build_data.py

Sources:
  - JTA Overnight Travel Statistics Survey (宿泊旅行統計調査), annual confirmed
    workbooks: https://www.mlit.go.jp/kankocho/tokei_hakusyo/shukuhakutokei.html
  - Prefecture outlines: https://github.com/dataofjapan/land (from MLIT 国土数値情報)
"""
import io
import json
import os
import re
import time

import numpy as np
import openpyxl
import pandas as pd
import requests

LATEST, BASE = 2025, 2019
URLS = {
    2025: "https://www.mlit.go.jp/kankocho/content/002010340.xlsx",
    2019: "https://www.mlit.go.jp/kankocho/tokei_hakusyo/content/001350484.xlsx",
}
BOUNDARIES_URL = "https://raw.githubusercontent.com/dataofjapan/land/master/japan.geojson"
MIN_NIGHTS = 500_000
METRO = {"Tokyo", "Kanagawa", "Chiba", "Saitama", "Aichi", "Osaka", "Kyoto", "Hyogo"}
LONG_HAUL = ["米国", "アメリカ", "カナダ", "英国", "イギリス", "ドイツ", "フランス", "イタリア", "スペイン", "オーストラリア"]
SKI_MONTHS = [12, 1, 2, 3]

REGIONS = {
    "Hokkaido": ["Hokkaido"],
    "Tohoku": ["Aomori", "Iwate", "Miyagi", "Akita", "Yamagata", "Fukushima"],
    "Kanto": ["Ibaraki", "Tochigi", "Gunma", "Saitama", "Chiba", "Tokyo", "Kanagawa"],
    "Chubu": ["Niigata", "Toyama", "Ishikawa", "Fukui", "Yamanashi", "Nagano", "Gifu", "Shizuoka", "Aichi"],
    "Kansai": ["Mie", "Shiga", "Kyoto", "Osaka", "Hyogo", "Nara", "Wakayama"],
    "Chugoku": ["Tottori", "Shimane", "Okayama", "Hiroshima", "Yamaguchi"],
    "Shikoku": ["Tokushima", "Kagawa", "Ehime", "Kochi"],
    "Kyushu & Okinawa": ["Fukuoka", "Saga", "Nagasaki", "Kumamoto", "Oita", "Miyazaki", "Kagoshima", "Okinawa"],
}
EN_TO_REGION = {p: r for r, ps in REGIONS.items() for p in ps}

PREFECTURES = [
    ("北海道", "Hokkaido", "Sapporo", 43.06, 141.35), ("青森県", "Aomori", "Aomori", 40.82, 140.74),
    ("岩手県", "Iwate", "Morioka", 39.70, 141.15), ("宮城県", "Miyagi", "Sendai", 38.27, 140.87),
    ("秋田県", "Akita", "Akita", 39.72, 140.10), ("山形県", "Yamagata", "Yamagata", 38.24, 140.36),
    ("福島県", "Fukushima", "Fukushima", 37.75, 140.47), ("茨城県", "Ibaraki", "Mito", 36.34, 140.45),
    ("栃木県", "Tochigi", "Utsunomiya", 36.57, 139.88), ("群馬県", "Gunma", "Maebashi", 36.39, 139.06),
    ("埼玉県", "Saitama", "Saitama", 35.86, 139.65), ("千葉県", "Chiba", "Chiba", 35.61, 140.12),
    ("東京都", "Tokyo", "Tokyo", 35.69, 139.69), ("神奈川県", "Kanagawa", "Yokohama", 35.45, 139.64),
    ("新潟県", "Niigata", "Niigata", 37.90, 139.02), ("富山県", "Toyama", "Toyama", 36.70, 137.21),
    ("石川県", "Ishikawa", "Kanazawa", 36.59, 136.63), ("福井県", "Fukui", "Fukui", 36.07, 136.22),
    ("山梨県", "Yamanashi", "Kofu", 35.66, 138.57), ("長野県", "Nagano", "Nagano", 36.65, 138.18),
    ("岐阜県", "Gifu", "Gifu", 35.39, 136.72), ("静岡県", "Shizuoka", "Shizuoka", 34.98, 138.38),
    ("愛知県", "Aichi", "Nagoya", 35.18, 136.91), ("三重県", "Mie", "Tsu", 34.73, 136.51),
    ("滋賀県", "Shiga", "Otsu", 35.00, 135.87), ("京都府", "Kyoto", "Kyoto", 35.02, 135.76),
    ("大阪府", "Osaka", "Osaka", 34.69, 135.52), ("兵庫県", "Hyogo", "Kobe", 34.69, 135.18),
    ("奈良県", "Nara", "Nara", 34.69, 135.83), ("和歌山県", "Wakayama", "Wakayama", 34.23, 135.17),
    ("鳥取県", "Tottori", "Tottori", 35.50, 134.24), ("島根県", "Shimane", "Matsue", 35.47, 133.05),
    ("岡山県", "Okayama", "Okayama", 34.66, 133.93), ("広島県", "Hiroshima", "Hiroshima", 34.40, 132.46),
    ("山口県", "Yamaguchi", "Yamaguchi", 34.19, 131.47), ("徳島県", "Tokushima", "Tokushima", 34.07, 134.56),
    ("香川県", "Kagawa", "Takamatsu", 34.34, 134.04), ("愛媛県", "Ehime", "Matsuyama", 33.84, 132.77),
    ("高知県", "Kochi", "Kochi", 33.56, 133.53), ("福岡県", "Fukuoka", "Fukuoka", 33.61, 130.42),
    ("佐賀県", "Saga", "Saga", 33.25, 130.30), ("長崎県", "Nagasaki", "Nagasaki", 32.74, 129.87),
    ("熊本県", "Kumamoto", "Kumamoto", 32.79, 130.74), ("大分県", "Oita", "Oita", 33.24, 131.61),
    ("宮崎県", "Miyazaki", "Miyazaki", 31.91, 131.42), ("鹿児島県", "Kagoshima", "Kagoshima", 31.56, 130.56),
    ("沖縄県", "Okinawa", "Naha", 26.21, 127.68),
]
JP_TO_EN = {p[0]: p[1] for p in PREFECTURES}


def num(v) -> float:
    try:
        return float(str(v).replace(",", "").lstrip("*"))
    except ValueError:
        return np.nan


def sum_columns(wb, sheet: str, keywords: list[str]) -> pd.Series:
    rows = [list(r) + [None] * 30 for r in wb[sheet].iter_rows(values_only=True)]
    header = [f"{rows[3][c] or ''}{rows[4][c] or ''}".replace("\n", "") for c in range(len(rows[3]))]
    cols = [c for c, h in enumerate(header) if any(k in h for k in keywords)]
    out = {}
    for r in rows:
        name = str(r[0] or "").strip("　 ").lstrip("0123456789")
        if name in JP_TO_EN:
            out[JP_TO_EN[name]] = sum(num(r[c]) for c in cols)
    return pd.Series(out)


def load_year(year: int) -> pd.DataFrame:
    print(f"downloading JTA workbook {year}…", flush=True)
    raw = requests.get(URLS[year], headers={"User-Agent": "Mozilla/5.0"}, timeout=180).content
    wb = openpyxl.load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    df = pd.DataFrame({
        "foreign_nights": sum_columns(wb, "第2表(年計)", ["外国人"]),
        "by_nationality": sum_columns(wb, "参考第1表(年計)", ["外国人"]),
        "long_haul": sum_columns(wb, "参考第1表(年計)", LONG_HAUL),
        "ski_occupancy": pd.concat(
            [sum_columns(wb, f"第8表({m}月)", ["リゾート"]) for m in SKI_MONTHS], axis=1).mean(axis=1),
    })
    # monthly baselines per prefecture, for comparing the live monthly release
    df["monthly_nights"] = pd.concat(
        [sum_columns(wb, f"第2表({m}月)", ["外国人"]) for m in range(1, 13)], axis=1).values.tolist()
    df["monthly_occupancy"] = pd.concat(
        [column_values(wb, f"第8表({m}月)", 1) for m in range(1, 13)], axis=1).values.tolist()
    return df


def column_values(wb, sheet: str, col: int) -> pd.Series:
    """One column by position, for each prefecture (used for the all-facility occupancy rate)."""
    out = {}
    for r in wb[sheet].iter_rows(values_only=True):
        name = str(r[0] or "").strip("　 ").lstrip("0123456789")
        if name in JP_TO_EN and len(r) > col:
            out[JP_TO_EN[name]] = num(r[col])
    return pd.Series(out)


# ---------------------------------------------------------------------------
# Miki Shoji Tokyo office history: each December page carries a 13-month table,
# so one page per year reconstructs the full monthly series.
# ---------------------------------------------------------------------------
def miki_tables(html: str) -> dict:
    strip = lambda s: re.sub(r"\s+", " ", re.sub(r"<[^>]+>", "", s)).strip()
    tables = {}
    for m in re.finditer(r"<h2[^>]*>(.*?)</h2>.*?<table>(.*?)</table>", html, re.S):
        title, body = strip(m.group(1)), m.group(2)
        ths = [strip(x) for x in re.findall(r"<th[^>]*>(.*?)</th>", body, re.S)]
        months, year = [], None
        for h in ths:
            full, part = re.match(r"^(\d{4})\.(\d{2})$", h), re.match(r"^(\d{2})$", h)
            if full:
                year = int(full.group(1)); months.append(f"{full.group(1)}-{full.group(2)}")
            elif part and year:
                if months and int(part.group(1)) < int(months[-1][5:]):
                    year += 1
                months.append(f"{year}-{part.group(1)}")
        rows = {}
        for tr in re.findall(r"<tr>(.*?)</tr>", body, re.S):
            tds = [strip(x) for x in re.findall(r"<td[^>]*>(.*?)</td>", tr, re.S)]
            if len(tds) > 1:
                rows[tds[0]] = [num(v) for v in tds[1:1 + len(months)]]
        if months:
            tables[title] = {"months": months, "rows": rows}
    return tables


def build_miki_history(first: int = 1991, last: int = LATEST) -> dict:
    series = {}
    for y in range(first, last + 1):
        url = f"https://www.e-miki.com/rent/tokyo.html?yyyy={y}&mm=12"
        try:
            html = requests.get(url, headers={"User-Agent": "Mozilla/5.0"}, timeout=60).text
            t = miki_tables(html)
            vac, rent = t["東京の平均空室率"], t["東京の平均賃料"]
            for i, mo in enumerate(vac["months"]):
                v, rr = vac["rows"]["平均空室率"][i], rent["rows"]["平均賃料"][i]
                if v is not None and rr is not None and not np.isnan(v) and not np.isnan(rr):
                    series[mo] = (v, rr)
        except Exception as e:  # a missing year just leaves a gap
            print(f"  miki {y}: {e}", flush=True)
        time.sleep(0.4)
    months = sorted(series)
    return {"months": months, "vacancy": [series[m][0] for m in months], "rent": [series[m][1] for m in months]}


def simplify_ring(ring: list) -> list:
    out = []
    for x, y in ring:
        p = [round(x, 2), round(y, 2)]
        if not out or out[-1] != p:
            out.append(p)
    return out


def build_boundaries() -> dict:
    print("downloading prefecture boundaries…", flush=True)
    geo = requests.get(BOUNDARIES_URL, timeout=180).json()
    for f in geo["features"]:
        en = JP_TO_EN[f["properties"]["nam_ja"]]
        f["id"] = en
        f["properties"] = {"name": en}  # ECharts matches regions by properties.name
        g = f["geometry"]
        polygons = g["coordinates"] if g["type"] == "MultiPolygon" else [g["coordinates"]]
        polygons = [[r for r in map(simplify_ring, p) if len(r) >= 4] for p in polygons]
        g["type"], g["coordinates"] = "MultiPolygon", [p for p in polygons if p]
    return geo


def clean_list(values, dp=0):
    if values is None:
        return None
    return [None if v is None or pd.isna(v) else (round(float(v), dp) if dp else round(float(v))) for v in values]


def main():
    out_dir = os.path.join(os.path.dirname(__file__), "..", "data")
    os.makedirs(out_dir, exist_ok=True)

    now, then = load_year(LATEST), load_year(BASE)

    df = pd.DataFrame(index=[p[1] for p in PREFECTURES])
    df["capital"] = [p[2] for p in PREFECTURES]
    df["lat"] = [p[3] for p in PREFECTURES]
    df["lon"] = [p[4] for p in PREFECTURES]
    df["region"] = [EN_TO_REGION[p[1]] for p in PREFECTURES]
    df["nights_latest"] = now["foreign_nights"]
    df["nights_base"] = then["foreign_nights"]
    df["growth_pct"] = 100 * (now["foreign_nights"] / then["foreign_nights"] - 1)
    df["ski_occupancy"] = now["ski_occupancy"]
    df["long_haul_share"] = 100 * now["long_haul"] / now["by_nationality"]

    # Destination screen: regional prefectures with enough demand; score = mean z of 3 metrics
    scored_cols = ["growth_pct", "ski_occupancy", "long_haul_share"]
    screen = df[~df.index.isin(METRO) & (df["nights_latest"] >= MIN_NIGHTS)].copy()
    z = (screen[scored_cols] - screen[scored_cols].mean()) / screen[scored_cols].std()
    screen["score"] = z.fillna(0).mean(axis=1)
    screen = screen.sort_values("score", ascending=False)
    ranks = {name: i + 1 for i, name in enumerate(screen.index)}
    scores = screen["score"].to_dict()

    prefs = []
    for name, r in df.iterrows():
        prefs.append({
            "name": name, "capital": r["capital"],
            "lat": r["lat"], "lon": r["lon"], "region": r["region"],
            "nights_latest": None if pd.isna(r["nights_latest"]) else round(float(r["nights_latest"])),
            "nights_base": None if pd.isna(r["nights_base"]) else round(float(r["nights_base"])),
            "growth_pct": None if pd.isna(r["growth_pct"]) else round(float(r["growth_pct"]), 1),
            "ski_occupancy": None if pd.isna(r["ski_occupancy"]) else round(float(r["ski_occupancy"]), 1),
            "long_haul_share": None if pd.isna(r["long_haul_share"]) else round(float(r["long_haul_share"]), 1),
            "monthly_nights_latest": clean_list(now["monthly_nights"].get(name)),
            "monthly_nights_base": clean_list(then["monthly_nights"].get(name)),
            "monthly_occupancy_latest": clean_list(now["monthly_occupancy"].get(name), 1),
            "screened": name in ranks,
            "rank": ranks.get(name),
            "score": round(float(scores[name]), 3) if name in scores else None,
        })

    payload = {
        "meta": {
            "latest_year": LATEST, "base_year": BASE, "min_nights": MIN_NIGHTS,
            "metro_excluded": sorted(METRO),
            "source": "Japan Tourism Agency, Overnight Travel Statistics Survey, annual confirmed values",
            "source_url": "https://www.mlit.go.jp/kankocho/tokei_hakusyo/shukuhakutokei.html",
        },
        "totals": {
            "foreign_nights_latest": round(float(now["foreign_nights"].sum())),
            "foreign_nights_base": round(float(then["foreign_nights"].sum())),
        },
        "prefectures": prefs,
    }

    with open(os.path.join(out_dir, "tourism.json"), "w", encoding="utf-8") as f:
        json.dump(payload, f, ensure_ascii=False)
    print(f"tourism.json written: total {LATEST} nights = {payload['totals']['foreign_nights_latest']:,} "
          f"(JTA published 2025 total: 179,922,130)", flush=True)

    geo = build_boundaries()
    with open(os.path.join(out_dir, "japan.geojson"), "w", encoding="utf-8") as f:
        json.dump(geo, f, ensure_ascii=False)
    print("japan.geojson written", flush=True)

    print("downloading Miki Shoji Tokyo office history…", flush=True)
    hist = build_miki_history()
    with open(os.path.join(out_dir, "miki_tokyo_history.json"), "w", encoding="utf-8") as f:
        json.dump(hist, f)
    print(f"miki_tokyo_history.json written: {hist['months'][0]} → {hist['months'][-1]} ({len(hist['months'])} months)", flush=True)


if __name__ == "__main__":
    main()
